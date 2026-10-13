"""v4.44 — Columna "Audio" al final del xlsx.

- `normalize_dossier_dataframe` genera "Audio" con la URL (Streaming - Imagen)
  solo para Tipo de Medio Radio/Televisión (Aire/AM/Cable/FM normalizan a esos);
  el resto de tipos queda vacío.
- `output_columns_for_export` y `_columnas_preservadas` ponen "Audio" como
  última columna en todos los modos, sin alterar el orden del resto.
- End-to-end: el xlsx final trae "Audio" al final con las URLs visibles.
"""
import io
import unittest

from openpyxl import Workbook, load_workbook

from pipeline import (
    _columnas_preservadas,
    load_dossier_dataframe,
    normalize_dossier_dataframe,
    output_columns_for_export,
    process_dossier,
)


def _dossier_audio():
    """Dossier sintético con Tipo de Medio y URL (Streaming - Imagen)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    headers = ["NoticiaId", "Fecha", "Medio", "Tipo de Medio", "Título",
               "CuerpoEs", "URL (Streaming - Imagen)", "URL Nota"]
    ws.append(headers)
    filas = [
        (1, "2026-10-01", "La FM", "FM", "Entrevista radial",
         "Resumen uno.", "https://stream.ejemplo.co/fm1", ""),
        (2, "2026-10-02", "Canal Aire", "Aire", "Noticiero",
         "Resumen dos.", "https://stream.ejemplo.co/tv1", ""),
        (3, "2026-10-03", "El Diario", "Prensa", "Nota escrita",
         "Resumen tres.", "https://stream.ejemplo.co/no-debe-salir", ""),
        (4, "2026-10-04", "Portal Web", "Internet", "Nota web",
         "Resumen cuatro.", "", "https://ejemplo.co/nota4"),
    ]
    for f in filas:
        ws.append(list(f))
    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


class TestColumnaAudioV444(unittest.TestCase):
    def test_audio_solo_radio_y_television(self):
        df = load_dossier_dataframe(_dossier_audio())
        dfn = normalize_dossier_dataframe(df, {}, {})
        por_titulo = dict(zip(dfn["Título"], dfn["Audio"]))
        self.assertEqual(por_titulo["Entrevista radial"],
                         "https://stream.ejemplo.co/fm1")
        self.assertEqual(por_titulo["Noticiero"],
                         "https://stream.ejemplo.co/tv1")
        self.assertEqual(por_titulo["Nota escrita"], "")
        self.assertEqual(por_titulo["Nota web"], "")

    def test_audio_al_final_en_todos_los_modos(self):
        self.assertEqual(output_columns_for_export()[-1], "Audio")
        self.assertEqual(
            output_columns_for_export(include_ai=True)[-1], "Audio")
        self.assertEqual(
            output_columns_for_export(include_ai=True,
                                      include_prominencia=True)[-2:],
            ["Prominencia", "Audio"])
        # El resto del orden no cambia: Audio solo se agrega al final.
        sin_audio = [c for c in output_columns_for_export(include_ai=True)
                     if c != "Audio"]
        self.assertEqual(sin_audio[-1], "Contexto analizado")

    def test_columnas_preservadas_audio_al_final(self):
        cols = _columnas_preservadas(["ID", "TÍTULO"], include_ai=True)
        self.assertEqual(cols[-1], "Audio")
        # No duplica si el xlsx de entrada ya trae la columna (se conserva
        # en su posición original, como el resto de columnas de entrada).
        cols2 = _columnas_preservadas(["ID", "Audio"], include_ai=True)
        self.assertEqual(cols2.count("Audio"), 1)
        self.assertEqual(cols2[1], "Audio")

    def test_e2e_audio_ultima_columna_con_urls(self):
        res = process_dossier(io.BytesIO(_dossier_audio()), {}, {},
                              ai_config=None)
        wb = load_workbook(io.BytesIO(res["output_data"]))
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers[-1], "Audio")
        idx_audio = len(headers) - 1
        idx_titulo = headers.index("Título")
        valores = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            valores[row[idx_titulo]] = row[idx_audio]
        self.assertEqual(valores["Entrevista radial"],
                         "https://stream.ejemplo.co/fm1")
        self.assertEqual(valores["Noticiero"],
                         "https://stream.ejemplo.co/tv1")
        self.assertIn(valores["Nota escrita"], (None, ""))
        self.assertIn(valores["Nota web"], (None, ""))


if __name__ == "__main__":
    unittest.main()

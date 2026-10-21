# -*- coding: utf-8 -*-
"""v4.46: los títulos que llegan con entidades HTML sin decodificar se
limpian a texto legible, sin quitar ni recortar nada más.

Pedido del usuario (2026-10-21): corregir títulos como
"&#x201C;La revaluaci&#xF3;n es la amenaza n&#xFA;mero uno...&#x201D;"
(decodificar las entidades) y NO quitar, eliminar ni recortar el título si
trae "|" o cosas similares: es solo un paso de limpieza legible del título.

Implementación (pipeline.py): `decode_html_entities` (html.unescape, hasta
3 pasadas por doble codificación) aplicada en `clean_title_for_output`,
que ahora también limpia la columna "Título" en
`normalize_dossier_dataframe`.
"""
from __future__ import annotations

import io
import os
import sys
import types
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Stub del paquete `openai` (mismo patrón que test_v432_pkl_solo.py).
_openai_stub = types.ModuleType('openai')
_openai_stub.OpenAI = object
sys.modules.setdefault('openai', _openai_stub)

from openpyxl import Workbook, load_workbook

from pipeline import (
    clean_title_for_output,
    decode_html_entities,
    process_dossier,
)

EJEMPLOS = [
    ("&#x201C;La revaluaci&#xF3;n es la amenaza n&#xFA;mero uno para los "
     "cafeteros&#x201D;: gremio habl&#xF3; de la ca&#xED;da del d&#xF3;lar",
     "“La revaluación es la amenaza número uno para los cafeteros”: "
     "gremio habló de la caída del dólar"),
    ("Expopet llegar&#xE1; a Medell&#xED;n en noviembre: fechas, "
     "actividades y m&#xE1;s informaci&#xF3;n",
     "Expopet llegará a Medellín en noviembre: fechas, actividades y "
     "más información"),
    ("Shell Rimula Fest 2026 convocar&#xE1; al gremio transportador en Guasca",
     "Shell Rimula Fest 2026 convocará al gremio transportador en Guasca"),
    ("Motos en Bogot&#xE1;: fechas y novedades de Expo 2 Ruedas 2026",
     "Motos en Bogotá: fechas y novedades de Expo 2 Ruedas 2026"),
    ("Mads Mikkelsen ser&#xE1; invitado estelar de Comic Con Colombia 2026",
     "Mads Mikkelsen será invitado estelar de Comic Con Colombia 2026"),
]


class TitulosEntidadesTests(unittest.TestCase):
    def test_ejemplos_del_usuario(self):
        for crudo, esperado in EJEMPLOS:
            with self.subTest(crudo=crudo):
                self.assertEqual(clean_title_for_output(crudo), esperado)

    def test_no_recorta_en_pipe_ni_en_guion(self):
        self.assertEqual(
            clean_title_for_output("Motos en Bogot&#xE1; | Expo 2 Ruedas 2026"),
            "Motos en Bogotá | Expo 2 Ruedas 2026",
        )
        self.assertEqual(
            clean_title_for_output("Nissan - Lanzamiento | An&#xE1;lisis completo"),
            "Nissan - Lanzamiento | Análisis completo",
        )

    def test_doble_codificacion(self):
        self.assertEqual(
            decode_html_entities("&amp;#x201C;Hola&amp;#x201D;"),
            "“Hola”",
        )

    def test_texto_plano_intacto(self):
        self.assertEqual(decode_html_entities("R&D | prueba - test"), "R&D | prueba - test")
        self.assertEqual(clean_title_for_output("Título normal sin entidades"),
                         "Título normal sin entidades")

    def test_proceso_completo_decodifica_el_titulo(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Título", "Resumen - Aclaracion", "Medio", "Fecha"])
        ws.append([EJEMPLOS[0][0], "resumen", "Prensa", "2026-09-01"])
        buf = io.BytesIO()
        wb.save(buf)

        res = process_dossier(buf.getvalue(), {}, {})
        out = load_workbook(io.BytesIO(res["output_data"]), read_only=True, data_only=True)
        sheet = out.active
        header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
        idx = header.index("Título")
        fila = next(sheet.iter_rows(min_row=2, max_row=2))
        self.assertEqual(fila[idx].value, EJEMPLOS[0][1])
        out.close()


if __name__ == "__main__":
    unittest.main()

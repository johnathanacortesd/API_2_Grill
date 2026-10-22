# -*- coding: utf-8 -*-
"""v4.47: el Resumen - Aclaracion (y el CuerpoEs) se mantienen exactos.

Bug reportado por el usuario (2026-10-22): la limpieza del resumen cortaba
palabras cuando el texto empezaba con un número. La causa era
`corregir_texto`: buscaba la primera letra mayúscula y descartaba todo lo
que hubiera antes ("2026 arranca con ventas récord en Colombia" salía
como "Colombia."), borraba los "[...]", colapsaba los párrafos y forzaba
un punto final; y en `normalize_dossier_dataframe` los medios gráficos
además rearmaban los párrafos con otro punto final forzado.

Regla del usuario: Resumen - Aclaracion y CuerpoEs deben mantenerse
exactos; la única limpieza permitida es la misma de los títulos (v4.46):
decodificar las entidades HTML ("creci&#xF3;" -> "creció"). Limpiar más
no eliminar ni quitar palabras. Aclaración del usuario (2026-10-22): el
texto debe terminar en punto, UN solo punto, no en tres ("..."): el
cierre lo garantiza `terminar_en_punto`, como ya hacía Grill (en
CardioAPI el cierre viejo agregaba "..." y eso era parte del bug).
"""
from __future__ import annotations

import io
import os
import sys
import types
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Stub del paquete `openai` (mismo patrón que test_v446_titulos_entidades.py).
_openai_stub = types.ModuleType('openai')
_openai_stub.OpenAI = object
sys.modules.setdefault('openai', _openai_stub)

from openpyxl import Workbook, load_workbook

from pipeline import (
    KEY_MAP,
    clean_cuerpo,
    corregir_texto,
    generate_output_excel,
    process_dossier,
)

EJEMPLO_USUARIO_CRUDO = (
    "Mercado automotor colombiano creci&#xF3; 36,1 % en septiembre "
    "y electrificados superan 50 %"
)
EJEMPLO_USUARIO_LIMPIO = (
    "Mercado automotor colombiano creció 36,1 % en septiembre "
    "y electrificados superan 50 %."
)


class CorregirTextoTests(unittest.TestCase):
    def test_empieza_con_numero_no_se_corta(self):
        self.assertEqual(
            corregir_texto("2026 arranca con ventas récord en Colombia"),
            "2026 arranca con ventas récord en Colombia.",
        )
        self.assertEqual(
            corregir_texto("36,1 % creció el mercado automotor en septiembre"),
            "36,1 % creció el mercado automotor en septiembre.",
        )
        self.assertEqual(
            corregir_texto("50 % de los carros vendidos ya son electrificados"),
            "50 % de los carros vendidos ya son electrificados.",
        )

    def test_empieza_en_minuscula_no_se_corta(self):
        self.assertEqual(
            corregir_texto("el mercado creció 36,1 % en septiembre, según Andemos"),
            "el mercado creció 36,1 % en septiembre, según Andemos.",
        )

    def test_ejemplo_del_usuario_decodifica_y_cierra_en_punto(self):
        # Decodifica la entidad, conserva los números y cierra en UN punto.
        self.assertEqual(corregir_texto(EJEMPLO_USUARIO_CRUDO), EJEMPLO_USUARIO_LIMPIO)

    def test_cierre_siempre_un_solo_punto_nunca_tres(self):
        self.assertEqual(corregir_texto("Texto sin punto final"), "Texto sin punto final.")
        self.assertEqual(corregir_texto("Texto que termina así..."), "Texto que termina así.")
        self.assertEqual(corregir_texto("Texto que termina así.."), "Texto que termina así.")
        self.assertEqual(corregir_texto("Texto que ya cierra bien."), "Texto que ya cierra bien.")

    def test_no_toca_corchetes_ni_parrafos_ni_elipsis_interna(self):
        self.assertEqual(
            corregir_texto("Primera parte [...] segunda parte"),
            "Primera parte [...] segunda parte.",
        )
        self.assertEqual(
            corregir_texto("Párrafo uno\n\nPárrafo dos sin punto"),
            "Párrafo uno\n\nPárrafo dos sin punto.",
        )
        self.assertEqual(corregir_texto("Uno... dos"), "Uno... dos.")

    def test_vacios(self):
        self.assertEqual(corregir_texto(""), "")
        self.assertEqual(corregir_texto("nan"), "")


class CleanCuerpoTests(unittest.TestCase):
    def test_decodifica_entidades_y_conserva_el_inicio(self):
        # clean_cuerpo (normalización) no toca el cierre: el punto final
        # único lo pone corregir_texto al exportar.
        self.assertEqual(
            clean_cuerpo(EJEMPLO_USUARIO_CRUDO),
            "Mercado automotor colombiano creció 36,1 % en septiembre "
            "y electrificados superan 50 %",
        )
        self.assertEqual(
            clean_cuerpo("2026 arranca con ventas r&#xE9;cord en Colombia"),
            "2026 arranca con ventas récord en Colombia",
        )

    def test_br_es_salto_de_linea_no_palabra(self):
        self.assertEqual(clean_cuerpo("Línea uno<br>Línea dos"), "Línea uno\nLínea dos")


class ProcesoCompletoTests(unittest.TestCase):
    def test_resumen_exacto_en_la_salida(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Título", "Resumen - Aclaracion", "Tipo de Medio", "Medio", "Fecha"])
        ws.append([
            "Ventas de carros en 2026",
            "2026 arranca con ventas r&#xE9;cord en Colombia",
            "Prensa", "El Tiempo", "2026-09-01",
        ])
        ws.append([
            "Mercado automotor en septiembre",
            EJEMPLO_USUARIO_CRUDO,
            "Internet", "Portafolio", "2026-09-01",
        ])
        buf = io.BytesIO()
        wb.save(buf)

        res = process_dossier(buf.getvalue(), {}, {})
        out = load_workbook(io.BytesIO(res["output_data"]), read_only=True, data_only=True)
        sheet = out.active
        filas = list(sheet.iter_rows(values_only=True))
        header = list(filas[0])
        idx_t = header.index("Título")
        idx_r = header.index("Resumen - Aclaracion")
        por_titulo = {f[idx_t]: f[idx_r] for f in filas[1:]}
        out.close()

        self.assertEqual(
            por_titulo["Ventas de carros en 2026"],
            "2026 arranca con ventas récord en Colombia.",
        )
        self.assertEqual(
            por_titulo["Mercado automotor en septiembre"],
            EJEMPLO_USUARIO_LIMPIO,
        )

    def test_cuerpoes_preservado_solo_decodifica(self):
        rows = [{
            "Título": "Mercado automotor en septiembre",
            "Resumen - Aclaracion": EJEMPLO_USUARIO_CRUDO,
            "CuerpoEs": "2026 arranca: el mercado creci&#xF3; 36,1 % en el pa&#xED;s",
        }]
        data = generate_output_excel(
            rows, KEY_MAP,
            columns_to_use=["Título", "Resumen - Aclaracion", "CuerpoEs"],
        )
        out = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        filas = list(out.active.iter_rows(values_only=True))
        out.close()
        self.assertEqual(filas[1][1], EJEMPLO_USUARIO_LIMPIO)
        self.assertEqual(
            filas[1][2],
            "2026 arranca: el mercado creció 36,1 % en el país.",
        )


if __name__ == "__main__":
    unittest.main()
